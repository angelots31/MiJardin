"""
Generación de documentos del quinto avance.

Todo lo que sale de la API como archivo se construye aquí:
  · factura_pdf()          -> factura de venta en PDF
  · reporte_ventas_pdf()   -> reporte diario de ventas en PDF
  · reporte_ventas_excel() -> el mismo reporte en .xlsx

Los documentos se arman en memoria (BytesIO) y se devuelven como bytes,
así que el servidor no necesita escribir nada en disco.
"""

import math
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.graphics.shapes import Circle, Drawing, Ellipse
from reportlab.platypus import (
    Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

# Paleta de la marca MiJardín (la misma del frontend)
VERDE = colors.HexColor("#23392E")
TERRACOTA = colors.HexColor("#D9714E")
SALVIA = colors.HexColor("#7C9473")
MOSTAZA = colors.HexColor("#E8AC4F")
CREMA = colors.HexColor("#FAF3E7")
GRIS = colors.HexColor("#E4DCCD")

EMPRESA = "MiJardín · Floristería digital"
NIT = "NIT 901.456.789-1"
CONTACTO = "Medellín, Antioquia · contacto@mijardin.com · +57 304 476 1403"


def _pesos(valor) -> str:
    """Formatea un número como pesos colombianos: $ 1.250.000"""
    try:
        return f"$ {float(valor):,.0f}".replace(",", ".")
    except (TypeError, ValueError):
        return "$ 0"


def _estilos():
    base = getSampleStyleSheet()
    return {
        "titulo": ParagraphStyle("titulo", parent=base["Title"], fontSize=20,
                                 textColor=VERDE, spaceAfter=2, alignment=0),
        "marca": ParagraphStyle("marca", parent=base["Normal"], fontSize=9.5,
                                textColor=SALVIA, spaceAfter=10),
        "seccion": ParagraphStyle("seccion", parent=base["Heading2"], fontSize=11.5,
                                  textColor=VERDE, spaceBefore=12, spaceAfter=6),
        "normal": ParagraphStyle("normal", parent=base["Normal"], fontSize=9.5, leading=13),
        "pie": ParagraphStyle("pie", parent=base["Normal"], fontSize=8,
                              textColor=colors.HexColor("#6B7B70"), alignment=TA_CENTER),
        "derecha": ParagraphStyle("derecha", parent=base["Normal"], fontSize=9.5, alignment=TA_RIGHT),
        "centro": ParagraphStyle("centro", parent=base["Normal"], fontSize=9.5, alignment=TA_CENTER),
    }


def logo_dibujo(size: float = 13 * mm) -> Drawing:
    """
    Dibuja el logo de MiJardín (flor estilizada con hojas) usando primitivas
    vectoriales de ReportLab.

    Se dibuja en código a propósito: así el PDF no depende de archivos
    externos ni de librerías extra como svglib, y sale igual de nítido a
    cualquier tamaño.
    """
    dibujo = Drawing(size, size)
    centro = size / 2.0

    # Dos hojas salvia bajo la flor
    for signo in (-1, 1):
        hoja = Ellipse(
            centro + signo * size * 0.19,
            centro - size * 0.35,
            size * 0.15,
            size * 0.075,
        )
        hoja.fillColor = SALVIA
        hoja.strokeColor = None
        dibujo.add(hoja)

    # Cinco pétalos terracota alrededor del centro
    for indice in range(5):
        angulo = math.radians(indice * 72)
        petalo = Ellipse(
            centro + math.sin(angulo) * size * 0.22,
            centro + math.cos(angulo) * size * 0.22,
            size * 0.15,
            size * 0.13,
        )
        petalo.fillColor = TERRACOTA
        petalo.strokeColor = None
        dibujo.add(petalo)

    # Centro mostaza
    nucleo = Circle(centro, centro, size * 0.14)
    nucleo.fillColor = MOSTAZA
    nucleo.strokeColor = None
    dibujo.add(nucleo)

    return dibujo


def _encabezado(story, estilos, titulo: str, subtitulo: str):
    """Encabezado con el logo a la izquierda y el título del documento."""
    texto = [
        Paragraph(titulo, estilos["titulo"]),
        Paragraph(f"{EMPRESA} · {NIT}<br/>{subtitulo}", estilos["marca"]),
    ]
    encabezado = Table([[logo_dibujo(), texto]], colWidths=[16 * mm, None])
    encabezado.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (0, 0), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(encabezado)


def _pie_generacion(story, estilos, usuario: str = ""):
    generado = datetime.now().strftime("%d/%m/%Y a las %I:%M %p")
    firma = f"Documento generado por el sistema MiJardín el {generado}."
    if usuario:
        firma += f" Usuario: {usuario}."
    story.append(Spacer(1, 14))
    story.append(Paragraph(firma, estilos["pie"]))
    story.append(Paragraph(CONTACTO, estilos["pie"]))


# ---------------------------------------------------------------------
# FACTURA DE VENTA EN PDF
# ---------------------------------------------------------------------
def factura_pdf(factura: dict) -> bytes:
    buffer = BytesIO()
    estilos = _estilos()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=16 * mm, bottomMargin=16 * mm,
        title=f"Factura {factura['numero_factura']}",
    )
    story = []
    _encabezado(
        story, estilos,
        f"Factura {factura['numero_factura']}",
        f"Emitida el {factura['fecha_emision']:%d/%m/%Y} · Estado: {factura['estado']}"
        if isinstance(factura.get("fecha_emision"), datetime)
        else f"Estado: {factura['estado']}",
    )

    # Datos del cliente
    datos = [
        ["Cliente", factura.get("cliente", "")],
        ["Documento", str(factura.get("cliente_documento", "") or "")],
        ["Correo", factura.get("cliente_email", "") or ""],
        ["Teléfono", str(factura.get("cliente_telefono", "") or "")],
        ["Dirección", factura.get("cliente_direccion", "") or ""],
        ["Venta asociada", factura.get("numero_venta", "") or ""],
    ]
    tabla_cliente = Table(datos, colWidths=[35 * mm, 135 * mm])
    tabla_cliente.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("TEXTCOLOR", (0, 0), (0, -1), SALVIA),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("LINEBELOW", (0, 0), (-1, -2), 0.3, GRIS),
    ]))
    story.append(tabla_cliente)

    # Detalle
    story.append(Paragraph("Detalle de la factura", estilos["seccion"]))
    filas = [["#", "Descripción", "Tipo", "Cant.", "Precio unitario", "Descuento", "Subtotal"]]
    for i, item in enumerate(factura.get("items", []), start=1):
        filas.append([
            str(i),
            item["nombre_item"],
            item.get("tipo_item", "producto").capitalize(),
            str(item["cantidad"]),
            _pesos(item["precio_unitario"]),
            _pesos(item.get("descuento", 0)),
            _pesos(item["subtotal"]),
        ])

    tabla = Table(filas, colWidths=[10 * mm, 60 * mm, 20 * mm, 15 * mm, 27 * mm, 20 * mm, 27 * mm], repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), VERDE),
        ("TEXTCOLOR", (0, 0), (-1, 0), CREMA),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, CREMA]),
        ("GRID", (0, 0), (-1, -1), 0.25, GRIS),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(tabla)

    # Totales
    totales = [
        ["Subtotal", _pesos(factura["subtotal"])],
        ["Descuento", f"- {_pesos(factura['descuento'])}"],
        ["IVA", _pesos(factura["impuestos"])],
        ["Total a pagar", _pesos(factura["total"])],
    ]
    tabla_totales = Table(totales, colWidths=[40 * mm, 40 * mm], hAlign="RIGHT")
    tabla_totales.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LINEABOVE", (0, -1), (-1, -1), 0.8, VERDE),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, -1), (-1, -1), TERRACOTA),
        ("FONTSIZE", (0, -1), (-1, -1), 12),
    ]))
    story.append(Spacer(1, 10))
    story.append(tabla_totales)

    if factura.get("observaciones"):
        story.append(Paragraph("Observaciones", estilos["seccion"]))
        story.append(Paragraph(str(factura["observaciones"]), estilos["normal"]))

    _pie_generacion(story, estilos)
    doc.build(story)
    return buffer.getvalue()


# ---------------------------------------------------------------------
# REPORTE DIARIO DE VENTAS EN PDF
# ---------------------------------------------------------------------
def reporte_ventas_pdf(ventas: list[dict], resumen: dict, rango: str, usuario: str = "") -> bytes:
    buffer = BytesIO()
    estilos = _estilos()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=14 * mm, rightMargin=14 * mm,
        topMargin=14 * mm, bottomMargin=14 * mm,
        title=f"Reporte de ventas {rango}",
    )
    story = []
    _encabezado(story, estilos, "Reporte de ventas", f"Periodo: {rango}")

    # Tarjetas de resumen convertidas en una fila de tabla
    tarjetas = [[
        f"Ventas registradas\n{resumen['cantidad']}",
        f"Total facturado\n{_pesos(resumen['total'])}",
        f"Ticket promedio\n{_pesos(resumen['promedio'])}",
        f"Unidades vendidas\n{resumen['unidades']}",
    ]]
    tabla_resumen = Table(tarjetas, colWidths=[62 * mm] * 4)
    tabla_resumen.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), CREMA),
        ("BOX", (0, 0), (-1, -1), 0.4, GRIS),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, GRIS),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TEXTCOLOR", (0, 0), (-1, -1), VERDE),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(tabla_resumen)

    story.append(Paragraph("Ventas del periodo", estilos["seccion"]))
    filas = [["Fecha", "N.º venta", "Cliente", "Productos y servicios", "Cant.", "Total", "Estado"]]
    for v in ventas:
        detalle = "; ".join(
            f"{i['nombre_item']} x{i['cantidad']}" for i in v.get("items", [])
        ) or "—"
        unidades = sum(int(i["cantidad"]) for i in v.get("items", []))
        fecha = v["fecha_venta"]
        fecha_txt = fecha.strftime("%d/%m/%Y %I:%M %p") if isinstance(fecha, datetime) else str(fecha)
        filas.append([
            fecha_txt,
            v["numero_venta"],
            v.get("cliente", ""),
            Paragraph(detalle, estilos["normal"]),
            str(unidades),
            _pesos(v["total"]),
            v["estado"],
        ])

    if len(filas) == 1:
        filas.append(["—", "—", "No hay ventas en este periodo.", "—", "—", "—", "—"])

    tabla = Table(
        filas,
        colWidths=[32 * mm, 26 * mm, 42 * mm, 90 * mm, 15 * mm, 28 * mm, 22 * mm],
        repeatRows=1,
    )
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), VERDE),
        ("TEXTCOLOR", (0, 0), (-1, 0), CREMA),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (4, 1), (5, -1), "RIGHT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, CREMA]),
        ("GRID", (0, 0), (-1, -1), 0.25, GRIS),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(tabla)

    total_general = Table(
        [["Total del periodo", _pesos(resumen["total"])]],
        colWidths=[45 * mm, 40 * mm], hAlign="RIGHT",
    )
    total_general.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 12),
        ("TEXTCOLOR", (0, 0), (-1, -1), TERRACOTA),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("LINEABOVE", (0, 0), (-1, 0), 0.8, VERDE),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(Spacer(1, 8))
    story.append(total_general)

    _pie_generacion(story, estilos, usuario)
    doc.build(story)
    return buffer.getvalue()


# ---------------------------------------------------------------------
# REPORTE DIARIO DE VENTAS EN EXCEL
# ---------------------------------------------------------------------
def reporte_ventas_excel(ventas: list[dict], resumen: dict, rango: str) -> bytes:
    wb = Workbook()

    verde = "FF23392E"
    crema = "FFFAF3E7"
    borde = Border(*[Side(style="thin", color="FFE4DCCD")] * 4)

    # --- Hoja 1: una fila por venta ---
    hoja = wb.active
    hoja.title = "Ventas"
    encabezados = [
        "Fecha", "N.º venta", "Cliente", "Correo", "Productos y servicios",
        "Unidades", "Subtotal", "Descuento", "IVA", "Total", "Método de pago",
        "Estado", "Factura",
    ]
    hoja.append([f"MiJardín · Reporte de ventas · {rango}"])
    hoja.append([])
    hoja.append(encabezados)

    hoja["A1"].font = Font(bold=True, size=14, color=verde)

    for celda in hoja[3]:
        celda.font = Font(bold=True, color=crema)
        celda.fill = PatternFill("solid", fgColor=verde)
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = borde

    for v in ventas:
        detalle = "; ".join(f"{i['nombre_item']} x{i['cantidad']}" for i in v.get("items", []))
        unidades = sum(int(i["cantidad"]) for i in v.get("items", []))
        fecha = v["fecha_venta"]
        hoja.append([
            fecha.strftime("%Y-%m-%d %H:%M") if isinstance(fecha, datetime) else str(fecha),
            v["numero_venta"],
            v.get("cliente", ""),
            v.get("cliente_email", ""),
            detalle,
            unidades,
            float(v["subtotal"]),
            float(v["descuento"]),
            float(v["impuestos"]),
            float(v["total"]),
            v.get("metodo_pago", ""),
            v["estado"],
            v.get("numero_factura") or "Sin factura",
        ])

    for fila in hoja.iter_rows(min_row=4, max_row=hoja.max_row):
        for celda in fila:
            celda.border = borde
        for col in ("G", "H", "I", "J"):
            hoja[f"{col}{fila[0].row}"].number_format = '"$" #,##0'

    anchos = [18, 15, 26, 28, 48, 10, 14, 12, 12, 14, 16, 12, 15]
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho

    # El filtro automático deja la tabla lista para analizar en Excel
    hoja.auto_filter.ref = f"A3:{get_column_letter(len(encabezados))}{max(hoja.max_row, 4)}"
    hoja.freeze_panes = "A4"

    # --- Hoja 2: resumen e indicadores ---
    resumen_hoja = wb.create_sheet("Resumen")
    resumen_hoja.append(["Indicador", "Valor"])
    resumen_hoja.append(["Periodo", rango])
    resumen_hoja.append(["Ventas registradas", resumen["cantidad"]])
    resumen_hoja.append(["Unidades vendidas", resumen["unidades"]])
    resumen_hoja.append(["Total facturado", float(resumen["total"])])
    resumen_hoja.append(["Ticket promedio", float(resumen["promedio"])])
    resumen_hoja.append(["Generado", datetime.now().strftime("%Y-%m-%d %H:%M")])

    for celda in resumen_hoja[1]:
        celda.font = Font(bold=True, color=crema)
        celda.fill = PatternFill("solid", fgColor=verde)
    resumen_hoja["B5"].number_format = '"$" #,##0'
    resumen_hoja["B6"].number_format = '"$" #,##0'
    resumen_hoja.column_dimensions["A"].width = 24
    resumen_hoja.column_dimensions["B"].width = 26

    # --- Hoja 3: detalle línea por línea (para tablas dinámicas) ---
    detalle_hoja = wb.create_sheet("Detalle")
    detalle_hoja.append(["N.º venta", "Fecha", "Cliente", "Tipo", "Ítem",
                         "Cantidad", "Precio unitario", "Descuento", "Subtotal"])
    for celda in detalle_hoja[1]:
        celda.font = Font(bold=True, color=crema)
        celda.fill = PatternFill("solid", fgColor=verde)
    for v in ventas:
        fecha = v["fecha_venta"]
        fecha_txt = fecha.strftime("%Y-%m-%d") if isinstance(fecha, datetime) else str(fecha)
        for item in v.get("items", []):
            detalle_hoja.append([
                v["numero_venta"], fecha_txt, v.get("cliente", ""),
                item.get("tipo_item", "producto"), item["nombre_item"],
                int(item["cantidad"]), float(item["precio_unitario"]),
                float(item.get("descuento", 0)), float(item["subtotal"]),
            ])
    for i, ancho in enumerate([15, 13, 26, 12, 40, 10, 16, 12, 14], start=1):
        detalle_hoja.column_dimensions[get_column_letter(i)].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
